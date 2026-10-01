import openpyxl

workbook=openpyxl.load_workbook(r"D:\Road Map\Python\Day 1\2\Employees.xlsx")
Input_Sheet=workbook["Sheet1"]
Output_Sheet=workbook.create_sheet("Output")

headers = [cell.value for cell in Input_Sheet[1]]

Output_Sheet["A1"]="Name"
Output_Sheet["B1"]="Department"
Output_Sheet["C1"]="Expense"
Output_Sheet["D1"]="Status"
Output_Sheet["E1"]="Classification"

Output_Row=2

for row in Input_Sheet.iter_rows(min_row=2, values_only=True):
    for header, value in zip(headers, row):
        print(f"{header} | {value}")
    Employee_Name=row[0]
    Employee_Department= row[1]
    Employee_Expense=row[2]
    Employee_Status=row[3]

    print("----------------------------")
    print("Employee Report")
    print("----------------------------")
    classification = []

    if Employee_Expense > 10000:
        classification.append("High Value")

    if Employee_Status == "Pending":
        classification.append("Requires Approval")

    if not classification:
        classification.append("Normal")

    classification_text = ", ".join(classification)


    Output_Sheet.cell(row=Output_Row,column=1).value=Employee_Name
    Output_Sheet.cell(row=Output_Row,column=2).value=Employee_Department
    Output_Sheet.cell(row=Output_Row,column=3).value=Employee_Expense
    Output_Sheet.cell(row=Output_Row,column=4).value=Employee_Status
    Output_Sheet.cell(row=Output_Row,column=5).value=classification_text

    Output_Row+=1

    workbook.save(r"D:\Road Map\Python\Day 1\2\Employees_Output.xlsx")


